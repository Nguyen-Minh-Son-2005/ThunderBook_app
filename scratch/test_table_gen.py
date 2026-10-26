import copy
import docx

doc_template = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')
tmpl_tbl = doc_template.tables[0]

doc_new = docx.Document()
# Add a copy of the table to doc_new
new_tbl_elm = copy.deepcopy(tmpl_tbl._tbl)
doc_new._body._body.append(new_tbl_elm)

# Now access the new table in doc_new
t = doc_new.tables[0]
print("New table rows:", len(t.rows))
print("Row 5 nested tables:", len(t.rows[5].cells[1].tables))
print("Row 6 nested tables:", len(t.rows[6].cells[1].tables))

# Test modifying cells
t.rows[0].cells[1].text = "UC1"
t.rows[0].cells[3].text = "Đăng ký tài khoản"
print("Row 0 modified:", t.rows[0].cells[1].text, "-", t.rows[0].cells[3].text)
