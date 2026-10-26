# -*- coding: utf-8 -*-
import sys
import os
import copy
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Import data
sys.path.append(r'D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch')
from use_cases_data import USE_CASES

def set_cell_text_and_font(cell, text, font_name="Times New Roman", font_size=Pt(12), bold=False):
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = font_name
            r.font.size = font_size
            r.font.bold = bold

def populate_nested_table(nt, data_rows, font_name="Times New Roman", font_size=Pt(11)):
    sample_tr = copy.deepcopy(nt.rows[1]._tr)
    while len(nt.rows) > 1:
        nt._tbl.remove(nt.rows[-1]._tr)
        
    for item in data_rows:
        new_tr = copy.deepcopy(sample_tr)
        nt._tbl.append(new_tr)
        last_row = nt.rows[-1]
        set_cell_text_and_font(last_row.cells[0], str(item[0]), font_name, font_size, bold=False)
        set_cell_text_and_font(last_row.cells[1], str(item[1]), font_name, font_size, bold=False)
        set_cell_text_and_font(last_row.cells[2], str(item[2]), font_name, font_size, bold=False)

def build_use_case_table(tmpl_tbl, uc_data):
    tbl_elm = copy.deepcopy(tmpl_tbl._tbl)
    doc_temp = docx.Document()
    doc_temp._body._body.append(tbl_elm)
    t = doc_temp.tables[0]
    
    fn = "Times New Roman"
    fz = Pt(12)
    
    # Row 0: UC Code and Name
    set_cell_text_and_font(t.rows[0].cells[1], uc_data['code'], fn, fz, bold=True)
    set_cell_text_and_font(t.rows[0].cells[3], uc_data['name'], fn, fz, bold=True)
    
    # Row 1: Actor
    set_cell_text_and_font(t.rows[1].cells[1], uc_data['actor'], fn, fz)
    
    # Row 2: Description
    set_cell_text_and_font(t.rows[2].cells[1], uc_data['description'], fn, fz)
    
    # Row 3: Trigger
    set_cell_text_and_font(t.rows[3].cells[1], uc_data['trigger'], fn, fz)
    
    # Row 4: Precondition
    set_cell_text_and_font(t.rows[4].cells[1], uc_data['precondition'], fn, fz)
    
    # Row 5: Main flow
    nt_main = t.rows[5].cells[1].tables[0]
    populate_nested_table(nt_main, uc_data['main_flow'], fn, Pt(11))
    
    # Row 6: Alternative flow
    nt_alt = t.rows[6].cells[1].tables[0]
    populate_nested_table(nt_alt, uc_data['alt_flow'], fn, Pt(11))
    
    # Row 7: Postcondition
    set_cell_text_and_font(t.rows[7].cells[1], uc_data['postcondition'], fn, fz)
    
    return t._tbl

def main():
    print("Loading template docx...")
    doc_template = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')
    tmpl_tbl = doc_template.tables[0]
    
    print("Loading main report docx...")
    doc_path = r'D:\HAUI_UNIVEERSITY\N04_UngDungBanSachThunderBook.docx'
    doc = docx.Document(doc_path)
    
    start_idx = None
    end_idx = None
    for i, p in enumerate(doc.paragraphs):
        if '3.2 Mô tả chi tiết các use case' in p.text:
            start_idx = i
        if '3.3 Phân tích các use case' in p.text:
            end_idx = i
            break
            
    print(f"Section 3.2 starts at P[{start_idx}], Section 3.3 starts at P[{end_idx}]")
    assert start_idx is not None and end_idx is not None
    
    target_p = doc.paragraphs[end_idx]
    
    print(f"Building and inserting {len(USE_CASES)} use case tables...")
    for idx, uc in enumerate(USE_CASES):
        print(f"  Inserting {uc['code']}: {uc['name']}...")
        
        # 1. Heading paragraph
        h = doc.add_paragraph(uc['title'], style='Heading 3')
        for r in h.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.bold = True
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        target_p._p.addprevious(h._p)
        
        # 2. Table
        tbl_elm = build_use_case_table(tmpl_tbl, uc)
        target_p._p.addprevious(tbl_elm)
        
        # 3. Spacing paragraph
        space_p = doc.add_paragraph()
        space_p.paragraph_format.space_before = Pt(0)
        space_p.paragraph_format.space_after = Pt(8)
        target_p._p.addprevious(space_p._p)
        
    print("Removing old paragraphs between 3.2 and 3.3...")
    old_paras = [doc.paragraphs[i] for i in range(start_idx + 1, end_idx)]
    print(f"  Found {len(old_paras)} old paragraphs to remove.")
    for p in old_paras:
        p._p.getparent().remove(p._p)
        
    print("Saving modified document...")
    doc.save(doc_path)
    print("Saved successfully to:", doc_path)

if __name__ == '__main__':
    main()
