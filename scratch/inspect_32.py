import docx

doc = docx.Document(r'D:\HAUI_UNIVEERSITY\N04_UngDungBanSachThunderBook.docx')

with open(r'D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\doc_section_32.txt', 'w', encoding='utf-8') as f:
    for i in range(130, min(265, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        f.write(f"P[{i}] (style={p.style.name}): {p.text}\n")
