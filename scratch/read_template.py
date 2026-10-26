import docx
import sys

doc = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')
with open(r'D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\template_content.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total tables: {len(doc.tables)}\n\n")
    
    f.write("--- PARAGRAPHS ---\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"[{i}] {p.text}\n")
            
    f.write("\n--- TABLES ---\n")
    for i, table in enumerate(doc.tables):
        f.write(f"\n=== TABLE {i} ({len(table.rows)} rows, {len(table.columns)} cols) ===\n")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f"Row {r_idx}: " + " | ".join(cells) + "\n")
