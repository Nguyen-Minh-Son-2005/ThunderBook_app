# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc_path = r'D:\HAUI_UNIVEERSITY\N04_UngDungBanSachThunderBook.docx'
    doc = docx.Document(doc_path)
    
    # Locate P[113] and P[114]
    p_23_idx = None
    p_c3_idx = None
    for i, p in enumerate(doc.paragraphs):
        if '2.3. Giới thiệu về kiến trúc MVVM tổng quan' in p.text or '2.3 Giới thiệu về kiến trúc' in p.text:
            p_23_idx = i
        if 'Chương 3. Phân tích và thiết kế hệ thống' in p.text:
            p_c3_idx = i
            break
            
    print(f"p_23_idx = {p_23_idx}, p_c3_idx = {p_c3_idx}")
    assert p_23_idx is not None and p_c3_idx is not None
    
    target_p = doc.paragraphs[p_c3_idx]
    
    def add_heading(text, level=2):
        style_name = f'Heading {level}'
        p = doc.add_paragraph(style=style_name)
        p.paragraph_format.space_before = Pt(12 if level==2 else 8)
        p.paragraph_format.space_after = Pt(6 if level==2 else 4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14 if level==2 else 13)
        run.font.bold = True
        target_p._p.addprevious(p._p)
        return p

    def add_para(runs_data, first_line_indent=Pt(18)):
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if first_line_indent:
            p.paragraph_format.first_line_indent = first_line_indent
        for item in runs_data:
            text = item[0]
            bold = item[1] if len(item) > 1 else False
            italic = item[2] if len(item) > 2 else False
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = bold
            run.font.italic = italic
        target_p._p.addprevious(p._p)
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-10)
        
        r_bullet = p.add_run("• ")
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(13)
        r_bullet.font.bold = True
        
        r_title = p.add_run(title + ": ")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(13)
        r_title.font.bold = True
        
        r_desc = p.add_run(desc)
        r_desc.font.name = 'Times New Roman'
        r_desc.font.size = Pt(13)
        
        target_p._p.addprevious(p._p)
        return p

    def add_image_with_caption(img_path, caption_text, width=Inches(5.6)):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        target_p._p.addprevious(p_img._p)
        
        p_cap = doc.add_paragraph(style='name_image')
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True
        run_cap.font.bold = False
        target_p._p.addprevious(p_cap._p)

    # 1. Main Heading
    add_heading("2.3 Giới thiệu về kiến trúc MVVM tổng quan", level=2)
    
    # 2. Section 2.3.1
    add_heading("2.3.1 Khái niệm và nguồn gốc kiến trúc MVVM", level=3)
    add_para([
        ("Mô hình kiến trúc ", False),
        ("MVVM (Model - View - ViewModel)", True),
        (" là mẫu kiến trúc phần mềm được các kỹ sư Microsoft giới thiệu vào năm 2005 nhằm tối ưu hóa việc phân tách giao diện người dùng (UI) khỏi tầng logic nghiệp vụ và dữ liệu. Trong những năm gần đây, Google đã chính thức chọn và khuyến nghị kiến trúc MVVM kết hợp cùng bộ thư viện ", False),
        ("Android Jetpack", True),
        (" trở thành tiêu chuẩn cốt lõi trong phát triển ứng dụng di động hiện đại (Modern Android Development - MAD).", False)
    ])
    add_para([
        ("Trước khi MVVM trở nên phổ biến, mô hình ", False),
        ("MVC (Model - View - Controller)", True),
        (" thường xuyên được áp dụng trong các ứng dụng Android. Tuy nhiên, trong môi trường Android, các lớp ", False),
        ("Activity", True),
        (" và ", False),
        ("Fragment", True),
        (" vừa đảm nhận hiển thị giao diện người dùng (View), vừa phải gánh vác xử lý sự kiện và điều phối dữ liệu (Controller). Điều này dẫn tới hiện tượng tiêu cực thường được gọi là ", False),
        ('"God Activity"', True, True),
        (" – các lớp giao diện phình to lên tới hàng nghìn dòng lệnh, phụ thuộc chặt chẽ vào nhau khiến việc bảo trì, sửa lỗi và nâng cấp trở nên vô cùng khó khăn. Đặc biệt, vòng đời (Lifecycle) phức tạp của Android với các sự kiện thay đổi cấu hình như xoay màn hình (Configuration Changes) thường xuyên dẫn đến việc tải lại dữ liệu không cần thiết, gây giật lag giao diện hoặc nghiêm trọng hơn là rò rỉ bộ nhớ (Memory Leak) và dừng ứng dụng đột ngột (Crash/ANR).", False)
    ])
    add_para([
        ("Kiến trúc MVVM giải quyết triệt để các hạn chế trên nhờ áp dụng nghiêm ngặt nguyên lý ", False),
        ("Phân tách mối quan tâm (Separation of Concerns)", True),
        (" và ", False),
        ("Kiến trúc hướng dữ liệu phản ứng (Data-Driven / Reactive Architecture)", True),
        (". Những ưu điểm vượt trội mà MVVM mang lại bao gồm:", False)
    ])
    add_bullet("Tách biệt rõ ràng tầng hiển thị và logic", "Tầng View chỉ chuyên biệt cho việc vẽ giao diện và tiếp nhận tương tác người dùng; toàn bộ logic tính toán và chuẩn bị dữ liệu được chuyển giao cho ViewModel.")
    add_bullet("Quản lý trạng thái theo vòng đời (Lifecycle-Aware)", "ViewModel được thiết kế để sống sót qua các sự kiện hủy và tạo lại của Activity khi xoay màn hình, bảo toàn dữ liệu giao diện mà không cần truy vấn lại từ cơ sở dữ liệu.")
    add_bullet("Khả năng kiểm thử độc lập (Testability cao)", "ViewModel không giữ bất kỳ tham chiếu nào tới Context hay Widget giao diện, cho phép lập trình viên dễ dàng viết Unit Test độc lập mà không phụ thuộc vào thiết bị thật hay máy ảo.")
    add_bullet("Tối ưu hóa làm việc nhóm (Maintainability & Scalability)", "Các thành viên trong nhóm có thể phát triển song song tầng giao diện (View), tầng điều phối (ViewModel) và tầng dữ liệu (Repository) mà không gây xung đột mã nguồn khi đẩy lên kho lưu trữ chung.")

    # 3. Section 2.3.2
    add_heading("2.3.2 Các thành phần cốt lõi trong kiến trúc MVVM", level=3)
    add_para([
        ("Mô hình kiến trúc MVVM chuẩn của Google Android bao gồm 5 thành phần phối hợp chặt chẽ với nhau:", False)
    ])
    add_bullet("View (Tầng giao diện người dùng)", "Gồm các Activity, Fragment và layout XML. View có nhiệm vụ hiển thị dữ liệu lên màn hình, lắng nghe các sự kiện thao tác của người dùng (nhấn nút, nhập văn bản, vuốt màn hình) và chuyển giao sự kiện đó cho ViewModel. View đóng vai trò là một Observer lắng nghe (observe) các luồng LiveData từ ViewModel để tự động cập nhật lại UI khi có dữ liệu mới. View tuyệt đối không gọi trực tiếp cơ sở dữ liệu hay thực thi tính toán nghiệp vụ.")
    add_bullet("ViewModel (Tầng điều phối và quản lý UI State)", "Kế thừa từ lớp androidx.lifecycle.ViewModel, đóng vai trò cầu nối trung gian giữa View và Repository. ViewModel nhận sự kiện từ View, gọi Repository để lấy hoặc ghi dữ liệu, sau đó đóng gói kết quả vào các biến LiveData để View quan sát. ViewModel hoàn toàn không giữ tham chiếu tới Activity/Fragment để triệt tiêu nguy cơ rò rỉ bộ nhớ.")
    add_bullet("LiveData & SingleLiveEvent (Cơ chế dữ liệu phản ứng)", "LiveData là lớp giữ dữ liệu có khả năng nhận biết vòng đời (Lifecycle-Aware). LiveData chỉ thông báo dữ liệu cho View khi View đang ở trạng thái hoạt động (STARTED hoặc RESUMED), tự động hủy theo dõi khi View bị hủy (DESTROYED). Trong dự án, SingleLiveEvent được bổ sung để xử lý các sự kiện kích hoạt một lần duy nhất (như hiển thị Toast thông báo thành công hoặc chuyển màn hình).")
    add_bullet("Repository Pattern (Tầng trừu tượng hóa dữ liệu)", "Đóng vai trò là Nguồn chân lý duy nhất (Single Source of Truth) cho toàn bộ ứng dụng. Repository che giấu chi tiết kỹ thuật của các nguồn dữ liệu bên dưới (Firebase Firestore, Firebase Storage, Firebase Auth, Google Gemini API). ViewModel chỉ tương tác với Repository thông qua các phương thức nghiệp vụ mức cao mà không cần biết cơ chế truy vấn cụ thể.")
    add_bullet("Model (Tầng thực thể dữ liệu)", "Bao gồm các lớp POJO (Plain Old Java Objects) định nghĩa cấu trúc của các đối tượng nghiệp vụ trong hệ thống như Book, User, Order, Category, Voucher, Review... Các lớp này ánh xạ trực tiếp với các Document được lưu trữ trên cơ sở dữ liệu Cloud Firestore.")

    # Image 1: Official Google Architecture Diagram
    add_image_with_caption(
        r'D:\HAUI_UNIVEERSITY\final-architecture.png',
        'Hình 2.1: Sơ đồ kiến trúc ứng dụng chuẩn khuyến nghị bởi Google Android (Nguồn: developer.android.com)',
        width=Inches(5.5)
    )

    # 4. Section 2.3.3
    add_heading("2.3.3 Luồng dữ liệu một chiều (Unidirectional Data Flow - UDF)", level=3)
    add_para([
        ("Kiến trúc MVVM trong Android vận hành dựa trên nguyên tắc ", False),
        ("Luồng dữ liệu một chiều (Unidirectional Data Flow - UDF)", True),
        (", đảm bảo tính minh bạch và kiểm soát chặt chẽ trạng thái ứng dụng:", False)
    ])
    add_bullet("Chiều sự kiện đi lên (Events Flow Up)", "Khi người dùng thực hiện một hành động (nhấn nút Mua hàng, gõ từ khóa tìm kiếm), View bắt sự kiện và gọi hàm tương ứng trên ViewModel. ViewModel tiếp tục chuyển tiếp yêu cầu xử lý sang Repository. Repository gọi các dịch vụ Firebase hoặc AI API bên dưới.")
    add_bullet("Chiều dữ liệu đi xuống (Data Flows Down)", "Khi cơ sở dữ liệu Firebase trả về kết quả (dưới dạng Snapshot thời gian thực hoặc Callback hoàn thành), Repository nhận dữ liệu, chuyển đổi thành các đối tượng Model và trả về cho ViewModel. ViewModel cập nhật dữ liệu vào các đối tượng LiveData. View đang quan sát LiveData sẽ tự động nhận được giá trị mới và kích hoạt cơ chế vẽ lại giao diện (Render UI).")

    # 5. Section 2.3.4
    add_heading("2.3.4 Mô hình triển khai kiến trúc MVVM trong dự án Thunder Book", level=3)
    add_para([
        ("Trong dự án ứng dụng bán sách Thunder Book, nhóm sinh viên đã triển khai đồng bộ và nghiêm ngặt kiến trúc MVVM trên toàn bộ mã nguồn Java của hệ thống. Cấu trúc thư mục dự án được tổ chức phân tầng rõ ràng:", False)
    ])
    add_bullet("Tầng View (com.example.bookapp.View)", "Được chia tách thành hai phân hệ độc lập: View/user (HomeActivity, HomeFragment, BookDetailActivity, CartFragment, CheckoutActivity, ChatBotActivity, WriteReviewActivity...) phục vụ khách hàng mua sắm; và View/admin (AdminDashboardActivity, ManageBookActivity, ManageCategoryActivity, StatisticActivity, AddEditBookActivity...) phục vụ quản trị viên.")
    add_bullet("Tầng ViewModel (com.example.bookapp.ViewModel)", "Gồm 34 lớp ViewModel chuyên biệt (HomeViewModel, CheckoutViewModel, ChatViewModel, WriteReviewViewModel, AdminStatisticViewModel...) phụ trách quản lý toàn bộ trạng thái giao diện và giữ an toàn dữ liệu qua các vòng đời màn hình.")
    add_bullet("Tầng Repository (com.example.bookapp.Repository)", "Gồm 23 lớp Repository (BookRepository, OrderRepository, UserRepository, ChatRepository, ReviewRepository, AdminBookRepository...) là nơi DUY NHẤT được phép truy xuất dữ liệu với Cloud Firestore, Firebase Storage và gọi Google Gemini API.")
    add_bullet("Tầng Model (com.example.bookapp.Model)", "Bao gồm các lớp thực thể dữ liệu ánh xạ trực tiếp từ Firestore như Book.java, User.java, Order.java, Category.java, Voucher.java, Review.java, Author.java, Publisher.java, ChatMessage.java.")

    # Image 2: ThunderBook Architecture Diagram
    add_image_with_caption(
        r'D:\HAUI_UNIVEERSITY\mvvm_architecture_thunderbook.png',
        'Hình 2.2: Mô hình triển khai kiến trúc MVVM thực tế trong ứng dụng Thunder Book',
        width=Inches(5.7)
    )
    
    # Finally, remove the old placeholder P[113]
    old_p = doc.paragraphs[p_23_idx]
    old_p._p.getparent().remove(old_p._p)
    print("Old placeholder P[113] removed.")
    
    doc.save(doc_path)
    print("Document saved successfully with Section 2.3 fully populated!")

if __name__ == '__main__':
    main()
