import docx

doc = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')
t = doc.tables[0]
print("Table style:", t.style.name if t.style else None)
for r_i in range(len(t.rows)):
    row = t.rows[r_i]
    print(f"Row {r_i}: {len(row.cells)} cells")
    for c_i, c in enumerate(row.cells):
        print(f"  c{c_i}: width={c.width}, text={repr(c.text[:30])}")
        if c.tables:
            nt = c.tables[0]
            print(f"    Nested table: {len(nt.rows)} rows, style={nt.style.name if nt.style else None}")
            for nr_i in range(min(2, len(nt.rows))):
                nr = nt.rows[nr_i]
                print(f"      nrow {nr_i}: {[nc.text.strip() for nc in nr.cells]}")
