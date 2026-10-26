import docx

doc = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')
t = doc.tables[0]

with open(r'D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\nested_details.txt', 'w', encoding='utf-8') as f:
    f.write(f"Table 0 rows: {len(t.rows)}\n")
    for r_idx in range(len(t.rows)):
        row = t.rows[r_idx]
        cells = row.cells
        f.write(f"\nRow {r_idx}: num cells={len(cells)}\n")
        seen = set()
        for c_idx, cell in enumerate(cells):
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            f.write(f"  Cell {c_idx} text len={len(cell.text)}, nested tables={len(cell.tables)}\n")
            if cell.tables:
                for nt_idx, nt in enumerate(cell.tables):
                    f.write(f"    Nested table {nt_idx}: {len(nt.rows)} rows, {len(nt.columns)} cols\n")
                    for nr_idx, nr in enumerate(nt.rows):
                        nr_texts = [nc.text.strip().replace('\n', ' ') for nc in nr.cells]
                        f.write(f"      NRow {nr_idx}: {' | '.join(nr_texts)}\n")
            else:
                f.write(f"    Text: {repr(cell.text)}\n")
