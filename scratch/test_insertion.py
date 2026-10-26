import docx

doc = docx.Document()
p1 = doc.add_paragraph("Heading 1")
p2 = doc.add_paragraph("Old paragraph 1")
p3 = doc.add_paragraph("Old paragraph 2")
p4 = doc.add_paragraph("Heading 2")

# We want to replace p2 and p3 with new content before p4
t = doc.add_table(rows=2, cols=2)
t.rows[0].cells[0].text = "Header"
# Move table before p4
p4._p.addprevious(t._tbl)

# Add new paragraph and move before p4
new_p = doc.add_paragraph("New inserted paragraph")
p4._p.addprevious(new_p._p)

# Remove p2 and p3
p2._p.getparent().remove(p2._p)
p3._p.getparent().remove(p3._p)

doc.save(r"D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\test_output.docx")

# Reopen and inspect
doc_check = docx.Document(r"D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\test_output.docx")
print("Paragraphs in saved doc:")
for p in doc_check.paragraphs:
    print(" ", repr(p.text))
print("Tables in saved doc:", len(doc_check.tables))
