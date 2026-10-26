import docx

doc = docx.Document(r'D:\HAUI_UNIVEERSITY\đặc tả use case.docx')

with open(r'D:\HAUI_UNIVEERSITY\AppMangaFGroup\scratch\sample_table.txt', 'w', encoding='utf-8') as f:
    for t_idx in range(min(3, len(doc.tables))):
        t = doc.tables[t_idx]
        f.write(f"\n==================== TABLE {t_idx} ====================\n")
        for r_idx, r in enumerate(t.rows):
            f.write(f"--- ROW {r_idx} ({len(r.cells)} cells) ---\n")
            seen_cells = set()
            for c_idx, c in enumerate(r.cells):
                # Avoid printing identical merged cells multiple times
                if c._tc in seen_cells:
                    continue
                seen_cells.add(c._tc)
                f.write(f"  [Cell {c_idx}]:\n{c.text}\n")
